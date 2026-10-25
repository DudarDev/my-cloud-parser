# -*- coding: utf-8 -*-
"""
Модуль многоформатного экспорта (CSV, XLSX, JSON, DOCX).
Обеспечивает профессиональное форматирование, стилизацию таблиц и включение сводной аналитики.
"""

import json
import logging
from datetime import datetime
from typing import List, Dict, Any, Optional
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')

def save_to_csv(data: List[Dict[str, Any]], filename: str, stats: Optional[Dict[str, Any]] = None):
    """Экспорт данных в формат CSV."""
    if not data:
        logging.warning("Нет данных для экспорта в CSV.")
        return
    try:
        df = pd.DataFrame(data)
        df.to_csv(filename, index=False, encoding='utf-8-sig')
        logging.info(f"Файл CSV успешно сохранен: {filename} ({len(data)} записей)")
    except Exception as e:
        logging.error(f"Ошибка при сохранении в CSV ({filename}): {e}")


def save_to_json(data: List[Dict[str, Any]], filename: str, stats: Optional[Dict[str, Any]] = None):
    """Экспорт данных в формат JSON со структурированными метаданными и статистикой."""
    if not data:
        logging.warning("Нет данных для экспорта в JSON.")
        return
    try:
        output_payload = {
            "meta": {
                "generated_at": datetime.now().isoformat(),
                "total_count": len(data)
            },
            "analytics": stats or {},
            "listings": data
        }
        with open(filename, 'w', encoding='utf-8') as f:
            json.dump(output_payload, f, ensure_ascii=False, indent=2)
        logging.info(f"Файл JSON успешно сохранен: {filename}")
    except Exception as e:
        logging.error(f"Ошибка при сохранении в JSON ({filename}): {e}")


def save_to_xlsx(data: List[Dict[str, Any]], filename: str, stats: Optional[Dict[str, Any]] = None):
    """
    Экспорт в Excel с профессиональной стилизацией:
    - Вкладка 1: Список объявлений с цветовым выделением заголовков, границами и автошириной колонок.
    - Вкладка 2: Аналитическая сводка с ключевыми показателями и распределением цен.
    """
    if not data:
        logging.warning("Нет данных для экспорта в XLSX.")
        return
    try:
        wb = openpyxl.Workbook()
        ws_listings = wb.active
        ws_listings.title = "Объявления"

        # Заголовки
        headers = [
            "Код", "Модель", "Память", "Цвет", "Состояние",
            "Цена", "Числовая цена ($)", "Батарея", "Оператор",
            "Продавец", "Локация", "Доставка", "Оценка выгоды"
        ]
        ws_listings.append(headers)

        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        for col_num in range(1, len(headers) + 1):
            cell = ws_listings.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Заполнение строк
        for row_idx, item in enumerate(data, start=2):
            row_data = [
                item.get('code', ''),
                item.get('model', ''),
                item.get('storage', ''),
                item.get('color', ''),
                item.get('condition', ''),
                item.get('price', ''),
                item.get('price_numeric', ''),
                item.get('battery', ''),
                item.get('carrier', ''),
                item.get('seller', ''),
                item.get('location', ''),
                item.get('shipping', ''),
                item.get('deal_score', '')
            ]
            ws_listings.append(row_data)

            # Чередование цвета строк для читаемости
            row_fill = PatternFill(start_color="F2F5F9" if row_idx % 2 == 0 else "FFFFFF", fill_type="solid")
            for col_idx in range(1, len(row_data) + 1):
                c = ws_listings.cell(row=row_idx, column=col_idx)
                c.fill = row_fill
                c.border = thin_border
                c.font = Font(name="Calibri", size=10)

        # Автоматическая ширина колонок
        for col in ws_listings.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_listings.column_dimensions[col_letter].width = max(max_len + 4, 12)

        # Вкладка 2: Аналитика
        if stats:
            ws_stats = wb.create_sheet(title="Аналитика")
            ws_stats.append(["Показатель", "Значение"])
            ws_stats.cell(row=1, column=1).font = Font(bold=True, color="FFFFFF")
            ws_stats.cell(row=1, column=1).fill = header_fill
            ws_stats.cell(row=1, column=2).font = Font(bold=True, color="FFFFFF")
            ws_stats.cell(row=1, column=2).fill = header_fill

            p = stats.get('price_stats', {})
            b = stats.get('battery_stats', {})

            stats_rows = [
                ["Всего объявлений", stats.get('total_listings', len(data))],
                ["Медианная цена", f"${p.get('median', 'N/A')}"],
                ["Минимальная цена", f"${p.get('min', 'N/A')}"],
                ["Максимальная цена", f"${p.get('max', 'N/A')}"],
                ["Средняя цена", f"${p.get('mean', 'N/A')}"],
                ["Средняя батарея", f"{b.get('average', 'N/A')}%"],
                ["Батарея ниже 80% (сервис)", f"{b.get('below_80', 0)} шт."]
            ]
            for sr in stats_rows:
                ws_stats.append(sr)

            ws_stats.column_dimensions['A'].width = 30
            ws_stats.column_dimensions['B'].width = 20

        wb.save(filename)
        logging.info(f"Файл Excel успешно сохранен: {filename}")
    except Exception as e:
        logging.error(f"Ошибка при сохранении в XLSX ({filename}): {e}")


def save_to_docx(data: List[Dict[str, Any]], filename: str, stats: Optional[Dict[str, Any]] = None):
    """
    Экспорт отчета в Microsoft Word с титульной частью, аналитическими сводками и карточками объявлений.
    """
    if not data:
        logging.warning("Нет данных для экспорта в DOCX.")
        return
    try:
        doc = Document()

        # Заголовок документа
        title = doc.add_heading("Отчет по рынку Swappa.com", level=1)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        subtitle = doc.add_paragraph(f"Дата генерации: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Всего объявлений: {len(data)}")
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Аналитическая секция
        if stats:
            doc.add_heading("1. Рыночная аналитика", level=2)
            p_stats = stats.get('price_stats', {})
            b_stats = stats.get('battery_stats', {})

            table = doc.add_table(rows=1, cols=4)
            table.style = 'Light Shading Accent 1'
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = "Медианная цена"
            hdr_cells[1].text = "Диапазон цен"
            hdr_cells[2].text = "Средняя батарея"
            hdr_cells[3].text = "Износ (<80%)"

            row_cells = table.add_row().cells
            row_cells[0].text = f"${p_stats.get('median', 'N/A')}"
            row_cells[1].text = f"${p_stats.get('min', 'N/A')} - ${p_stats.get('max', 'N/A')}"
            row_cells[2].text = f"{b_stats.get('average', 'N/A')}%"
            row_cells[3].text = f"{b_stats.get('below_80', 0)} шт."

        # Секция объявлений
        doc.add_heading("2. Детальный список объявлений", level=2)
        for idx, item in enumerate(data, 1):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)

            r_title = p.add_run(f"#{idx} | {item.get('model', 'Phone')} {item.get('storage', '')} - {item.get('price', 'N/A')}")
            r_title.bold = True
            r_title.font.size = Pt(12)

            info_p = doc.add_paragraph()
            info_p.paragraph_format.left_indent = Inches(0.2)
            info_p.add_run(f"• Цвет: {item.get('color', 'N/A')} | Состояние: {item.get('condition', 'N/A')} | Батарея: {item.get('battery', 'N/A')}\n")
            info_p.add_run(f"• Оператор: {item.get('carrier', 'N/A')} | Доставка: {item.get('shipping', 'N/A')}\n")
            info_p.add_run(f"• Продавец: {item.get('seller', 'N/A')} ({item.get('location', 'N/A')}) | Код: {item.get('code', 'N/A')}")

        doc.save(filename)
        logging.info(f"Файл Word успешно сохранен: {filename}")
    except Exception as e:
        logging.error(f"Ошибка при сохранении в DOCX ({filename}): {e}")
